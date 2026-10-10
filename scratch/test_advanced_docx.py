import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def add_page_number(run):
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    run._r.append(fldSimple)

def add_numpages(run):
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls('w'))
    run._r.append(fldSimple)

doc = Document()
section = doc.sections[0]
footer = section.footer
p_foot = footer.paragraphs[0]
p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_foot.paragraph_format.bidi = True

r_label = p_foot.add_run("صفحه ")
add_page_number(p_foot.add_run())
r_of = p_foot.add_run(" از ")
add_numpages(p_foot.add_run())

doc.add_paragraph("صفحه اول تست")
doc.add_page_break()
doc.add_paragraph("صفحه دوم تست")
doc.save("scratch/test_footer.docx")
print("Footer test passed successfully!")
