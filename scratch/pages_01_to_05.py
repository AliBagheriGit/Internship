# -*- coding: utf-8 -*-
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches
from report_helpers import (
    apply_run_font, apply_latin_font, add_p, add_heading_1, add_heading_2, 
    add_heading_3, add_bullet, add_callout, add_code_box, add_table_data
)

def build_page_01(doc):
    """صفحه ۱: صفحه جلد رسمی دانشگاهی"""
    # Header space
    p_top = add_p(doc, "", space_after=10)
    
    # University Header
    p_uni = add_p(doc, "جمهوری اسلامی ایران", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=11, bold_prefix="", space_after=2)
    p_azad = add_p(doc, "دانشگاه آزاد اسلامی", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=14, bold_prefix="", space_after=2, color_rgb=(15, 53, 98))
    p_parand = add_p(doc, "واحد پرند", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=15, bold_prefix="", space_after=2, color_rgb=(15, 53, 98))
    p_faculty = add_p(doc, "دانشکده فنی و مهندسی — گروه آموزشی مهندسی کامپیوتر", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold_prefix="", space_after=14, color_rgb=(13, 92, 117))
    
    # Decorative line
    add_callout(doc, "گزارش نهایی دوره کارآموزی مهندسی کامپیوتر (مقطع کارشناسی)", [
        "رشته مهندسی کامپیوتر — نیم‌سال تحصیلی ۱۴۰۵-۱۴۰۴",
        "محل کارآموزی: استارتاپ استودیو و شتاب‌دهنده سلامت هوشمند «شریف کلینیک»"
    ], bg_hex="F0F9FF", border_color="0284C7")
    
    # Main Report Title
    p_rep = add_p(doc, "گزارش فنی و تخصصی کارآموزی", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=13, bold_prefix="", space_after=4, space_before=6)
    p_title = add_p(doc, "طراحی، معماری، پیاده‌سازی و بین‌المللی‌سازی پلتفرم وب شتاب‌دهنده و استارتاپ استودیوی سلامت هوشمند «شریف کلینیک»", 
                    align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=14.5, bold_prefix="", space_after=8, color_rgb=(15, 53, 98), line_spacing=1.25)
    
    p_tech = add_p(doc, "بر پایه پشته مدرن وب: Next.js 16 (App Router)، React 19، TypeScript 5، Tailwind CSS v4، مدیریت وضعیت Zustand و استقرار ابری Cloudflare Pages",
                   align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=10, bold_prefix="", space_after=14, color_rgb=(100, 116, 139))
    
    # Details Table
    details_data = [
        ["نام و نام خانوادگی دانشجو:", "علی باقری کزازی"],
        ["شماره دانشجویی:", "962070856"],
        ["رشته و مقطع تحصیلی:", "مهندسی کامپیوتر (مقطع کارشناسی) — دانشگاه آزاد اسلامی واحد پرند"],
        ["استاد محترم راهنما و ارزیاب:", "سرکار خانم فاطمه وردی"],
        ["سرپرست کارآموزی در شرکت:", "جناب آقای مهندس الیاس اوستاد"],
        ["عنوان و نقش در پروژه:", "طراح رابط کاربری و توسعه‌دهنده فرانت‌اند (Senior Frontend Developer)"],
        ["نام محل کارآموزی:", "شتاب‌دهنده و استارتاپ استودیوی سلامت هوشمند شریف کلینیک (Sharif Clinic)"],
        ["نشانی محل کارآموزی:", "تهران، ضلع غربی دانشگاه صنعتی شریف، خیابان شهید صادقی، پلاک ۲۶، واحد ۱۰"],
        ["مخزن آنلاین پروژه در گیت‌هاب:", "https://github.com/AliBagheriGit/Internship"],
        ["تاریخ ارائه گزارش:", "نیم‌سال تحصیلی ۱۴۰۵ - ۱۴۰۴"]
    ]
    add_table_data(doc, ["عنوان مشخصه", "اطلاعات ثبت‌شده"], details_data, caption="")

def build_page_02(doc):
    """صفحه ۲: صفحه تقدیم و سپاسگزاری"""
    add_heading_1(doc, "صفحه تقدیم و سپاسگزاری", space_before=10, space_after=12)
    
    add_heading_2(doc, "تقدیم به:", space_before=6, space_after=4)
    add_p(doc, "به نام یگانه هستی‌بخش و دانای مطلق،", align=WD_ALIGN_PARAGRAPH.RIGHT, size_pt=11, bold_prefix="", space_after=6)
    add_p(doc, "این اثر علمی، پژوهشی و مهندسی را با کمال احترام و تواضع تقدیم می‌دارم به:", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=10.5, space_after=4)
    add_bullet(doc, "پدر و مادر بزرگوار، صبور و فداکارم که زحمات بی‌دریغ، دعای خیر و مهربانی بی‌پایانشان همواره تکیه‌گاه، راهنما و چراغ روشن مسیر دانش و اخلاق من بوده است.", bold_prefix="خانواده دلسوز و فداکارم: ")
    add_bullet(doc, "که همواره در جهت اعتلای سلامت جامعه، نجات جان انسان‌ها و پیوند فناوری و طبابت مجاهدت و تلاش می‌نمایند.", bold_prefix="پزشکان فداکار، کادر زحمتکش درمان و مدافعان سلامت کشور: ")
    add_bullet(doc, "که با تکیه بر تفکر تحلیلی و تخصص نرم‌افزاری، افق‌های نوینی را در تحول دیجیتال و سلامت هوشمند برای آینده ایران رقم می‌زنند.", bold_prefix="پژوهشگران و مهندسان جوان حوزه سلامت دیجیتال: ")
    
    add_heading_2(doc, "سپاسگزاری و قدردانی ویژه:", space_before=10, space_after=4)
    add_p(doc, "بر خود فرض می‌دانم مراتب امتنان و سپاسگزاری صمیمانه خود را تقدیم نمایم به:", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=10.5, space_after=4)
    add_bullet(doc, "سرکار خانم فاطمه وردی، استاد محترم و فرزانه درس کارآموزی در دانشگاه آزاد اسلامی واحد پرند، که با راهنمایی‌های دلسوزانه، نظرات تخصصی و هدایت‌های علمی ارزشمندشان، نقشی محوری در ارتقای کیفیت آکادمیک و غنای مهندسی این گزارش ایفا نمودند.", bold_prefix="استاد محترم راهنما، سرکار خانم فاطمه وردی: ")
    add_bullet(doc, "جناب آقای مهندس الیاس اوستاد، سرپرست محترم کارآموزی در مجموعه شریف کلینیک، که با فراهم آوردن بستری پویا، خلاق و مبتنی بر همکاری فنی و با انتقال تجربیات گرانبهای صنعتی، زمینه رشد مهارت‌های عملی اینجانب را مهیا ساختند.", bold_prefix="سرپرست محترم کارآموزی، جناب آقای الیاس اوستاد: ")
    add_bullet(doc, "مدیران ارشد، پزشکان مشاور، توسعه‌دهندگان و پرسنل گرامی شتاب‌دهنده و استارتاپ استودیوی شریف کلینیک که با همکاری صمیمانه در محیط ناحیه نوآوری شریف، ما را در پیشبرد این پلتفرم یاری کردند.", bold_prefix="تیم تخصصی و کادر محترم شریف کلینیک: ")
    
    add_callout(doc, "بیانیه تعهد مهندسی در نظام سلامت", [
        "«مهندسی نرم‌افزار در ساحت سلامت و بهداشت، فراتر از کدنویسی صرف، مسئولیتی مستقیم در قبال حیات و کیفیت زندگی انسان‌هاست. این پروژه گامی استوار در جهت تجلی اخلاق حرفه‌ای، دقت مهندسی و ارزش‌آفرینی در نظام سلامت مدرن به شمار می‌رود.»"
    ], bg_hex="F8FAFC", border_color="1E3A8A")

def build_page_03(doc):
    """صفحه ۳: چکیده تفصیلی و کلمات کلیدی (فارسی و انگلیسی)"""
    add_heading_1(doc, "چکیده تفصیلی پروژه (Executive Abstract)", space_before=6, space_after=6)
    
    add_heading_2(doc, "چکیده فارسی:", space_before=4, space_after=3)
    add_p(doc, "پروژه حاضر حاصل دوره کارآموزی مهندسی کامپیوتر توسط دانشجو علی باقری کزازی (به شماره دانشجویی ۹۶۲۰۷۰۸۵۶) در دانشگاه آزاد اسلامی واحد پرند تحت نظارت استاد راهنما سرکار خانم فاطمه وردی و سرپرستی جناب آقای الیاس اوستاد در مجموعه شریف کلینیک به انجام رسیده است. هدف بنیادین پروژه، طراحی، معماری، پیاده‌سازی و بین‌المللی‌سازی پلتفرم وب رسمی «شریف کلینیک» (استارتاپ استودیو و شتاب‌دهنده تخصصی سلامت هوشمند واقع در ناحیه نوآوری شریف) بوده است تا پل ارتباطی پایداری میان تخصص مهندسی نرم‌افزار و نیازهای بالینی پزشکی برقرار سازد.", 
          space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_p(doc, "در لایه معماری نرم‌افزار، از جدیدترین پشته وب شامل فریم‌ورک Next.js 16 مبتنی بر App Router، کتابخانه React 19، زبان تایپ‌اسکریپت ۵ در حالت Strict، فریم‌ورک استایل‌دهی Tailwind CSS v4 و مدیریت وضعیت متمرکز Zustand استفاده شده است. سامانه مشتمل بر ویترین پویای شاخص‌های سلامت، کاتالوگ ۸ گانه خدمات بالینی با قابلیت فیلترینگ چندبعدی، مسیرهای دوگانه همکاری، ارکان چهارگانه جذب سرمایه، فرم هوشمند اعتبارسنجی ثبت طرح و پشتیبانی بین‌المللی کامل از دو زبان فارسی (RTL) و انگلیسی (LTR) است. همچنین فرآیند بیلد با رفع باگ توربوپک در فونت‌های متغیر و فعال‌سازی خروجی استاتیک (Static Export) برای استقرار موفق روی شبکه لبه Cloudflare Pages بهینه‌سازی گردید.", 
          space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_p(doc, "واژگان کلیدی: سلامت هوشمند، استارتاپ استودیو، شتاب‌دهنده سلامت، نکس‌جی‌اس ۱۶، ری‌اکت ۱۹، تایپ‌اسکریپت، زوستند، دوزبانه (i18n)، کلودفلر، دانشگاه آزاد اسلامی واحد پرند، شریف کلینیک.", 
          bold_prefix="واژگان کلیدی (Keywords): ", size_pt=10, color_rgb=(13, 92, 117), space_after=8)
    
    add_heading_2(doc, "Abstract (English):", space_before=6, space_after=3)
    p_en = doc.add_paragraph()
    p_en.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_en.paragraph_format.bidi = False
    p_en.paragraph_format.space_after = Pt(4)
    p_en.paragraph_format.line_spacing = 1.15
    r_en = p_en.add_run(
        "This internship report presents the end-to-end design, system architecture, internationalization, and cloud deployment of the official web platform for 'Sharif Clinic'—a premier smart health venture studio and accelerator located in the Sharif Innovation District. Completed by Ali Bagheri Kazazi (Student ID: 962070856) at Islamic Azad University - Parand Branch, under the academic guidance of Prof. Fatemeh Verdi and industry mentorship of Eng. Elyas Oustad, the platform bridges cutting-edge software engineering with clinical healthcare delivery. The technology stack leverages Next.js 16 (App Router), React 19, TypeScript 5, Tailwind CSS v4, and Zustand state store. Key achievements encompass an interactive 8-service healthcare catalog with category filtering, dual collaboration pathway models (Venture Studio vs. Accelerator), global digital health statistical dashboard, pitch submission modal with robust client-side validation, full bilingual support (Persian RTL & English LTR), and static export deployment optimization for Cloudflare Pages. Performance audits demonstrate near-perfect Google Lighthouse scores and superior responsiveness across all device form factors."
    )
    apply_latin_font(r_en, font_name="Calibri", size_pt=9.5, color_rgb=(30, 41, 59))
    
    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_kw.paragraph_format.bidi = False
    p_kw.paragraph_format.space_after = Pt(4)
    r_kw_title = p_kw.add_run("Keywords: ")
    apply_latin_font(r_kw_title, font_name="Calibri", size_pt=9.5, bold=True, color_rgb=(15, 53, 98))
    r_kw = p_kw.add_run("Smart Healthcare, Venture Studio, Accelerator, Next.js 16, React 19, TypeScript, Zustand, Internationalization (i18n), Cloudflare Pages, Parand Islamic Azad University, Sharif Clinic.")
    apply_latin_font(r_kw, font_name="Calibri", size_pt=9.5, color_rgb=(13, 92, 117))

def build_page_04(doc):
    """صفحه ۴: فهرست مطالب و جداول"""
    add_heading_1(doc, "فهرست مطالب و ساختار گزارش", space_before=6, space_after=8)
    
    toc_data = [
        ["عنوان بخش / فصل", "شماره صفحه"],
        ["صفحه جلد رسمی، تقدیم، چکیده تفصیلی و فهرست مطالب", "صفحه ۱ تا ۴"],
        ["فصل اول: معرفی محل کارآموزی، زیست‌بوم و اهداف دوره کارآموزی", "صفحه ۵ تا ۷"],
        ["  • ۱.۱ تا ۱.۴: پیشینه، مأموریت و زیست‌بوم ناحیه نوآوری شریف", "صفحه ۵"],
        ["  • ۱.۵ تا ۱.۷: ساختار سازمانی، مدل‌های همکاری و شبکه شرکای درمانی", "صفحه ۶"],
        ["  • ۱.۸ تا ۱.۱۰: جایگاه کارآموز، شرح وظایف و گانت‌چارت هفتگی پیشبرد پروژه", "صفحه ۷"],
        ["فصل دوم: مفاهیم پایه، تحلیل مسئله و نیازمندی‌های سامانه سلامت هوشمند", "صفحه ۸ تا ۱۱"],
        ["  • ۲.۱ تا ۲.۳: مفاهیم پایه سلامت هوشمند و چالش‌های انفورماتیک سلامت", "صفحه ۸"],
        ["  • ۲.۴ تا ۲.۵: تحلیل ذی‌نفعان، پرسونای مخاطبان و بنچ‌مارک الگوهای جهانی", "صفحه ۹"],
        ["  • ۲.۶: تحلیل جامع نیازمندی‌های کارکردی (FR-01 تا FR-10) و ماتریس ردیابی", "صفحه ۱۰"],
        ["  • ۲.۷ تا ۲.۸: نیازمندی‌های غیرکارکردی (NFRs) و دیاگرام جریان کاربر", "صفحه ۱۱"],
        ["فصل سوم: معماری نرم‌افزار، زیرساخت و پشته فناوری نوین", "صفحه ۱۲ تا ۱۶"],
        ["  • ۳.۱ تا ۳.۳: معماری کلان سیستم و تحلیل عمیق Next.js 16 App Router", "صفحه ۱۲"],
        ["  • ۳.۴ تا ۳.۶: تحلیل React 19، کامپوننت‌های سرور/کلاینت و تایپ‌اسکریپت ۵", "صفحه ۱۳"],
        ["  • ۳.۷ تا ۳.۹: سیستم استایل‌دهی Tailwind CSS v4، زبان بصری و فونت وزیرمتن", "صفحه ۱۴"],
        ["  • ۳.۱۰ تا ۳.۱۲: معماری مدیریت وضعیت Zustand در برابر Redux و Context", "صفحه ۱۵"],
        ["  • ۳.۱۳ تا ۳.۱۵: معماری سئو پیشرفته، OpenGraph و بهینه‌سازی Web Vitals", "صفحه ۱۶"],
        ["فصل چهارم: طراحی تفصیلی و پیاده‌سازی ماژول‌های سامانه", "صفحه ۱۷ تا ۲۳"],
        ["  • ۴.۱ تا ۴.۳: ساختار پوشه‌ها، چیدمان ریشه (Layout) و خطایابی سراسری", "صفحه ۱۷"],
        ["  • ۴.۴ تا ۴.۵: پیاده‌سازی هدر هوشمند، سوئیچر زبان (Language Switcher) و منوی موبایل", "صفحه ۱۸"],
        ["  • ۴.۶ تا ۴.۷: پیاده‌سازی ویترین اصلی (Hero) و نشانگر زنده سلامت", "صفحه ۱۹"],
        ["  • ۴.۸ تا ۴.۱۰: پیاده‌سازی کاتالوگ تعاملی ۸ خدمت و فیلترینگ دسته‌بندی", "صفحه ۲۰"],
        ["  • ۴.۱۱ تا ۴.۱۲: پیاده‌سازی سیستم مودال سراسری و جزئیات خدمت بالینی", "صفحه ۲۱"],
        ["  • ۴.۱۳ تا ۴.۱۵: ماژول مسیرهای دوگانه همکاری، سرمایه‌گذاری و آمار جهانی", "صفحه ۲۲"],
        ["  • ۴.۱۶ تا ۴.۱۸: فرم ثبت طرح، سوالات متداول (FAQ) و فوتر چندستونه", "صفحه ۲۳"],
        ["فصل پنجم: ارزیابی کیفی، آزمون، بهینه‌سازی کارایی و استقرار ابری", "صفحه ۲۴ تا ۲۶"],
        ["  • ۵.۱ تا ۵.۳: استراتژی آزمون و ارزیابی استاتیک با TypeScript و ESLint", "صفحه ۲۴"],
        ["  • ۵.۴ تا ۵.۶: فرآیند بیلد Turbopack، رفع باگ فونت و استقرار Cloudflare Pages", "صفحه ۲۵"],
        ["  • ۵.۷ تا ۵.۹: تست واکنش‌گرایی، امنیت فرانت‌اند و خط‌لوله استقرار CI/CD", "صفحه ۲۶"],
        ["فصل ششم: تحلیل نتایج، دستاوردها و چشم‌انداز آینده", "صفحه ۲۷ تا ۲۸"],
        ["  • ۶.۱ تا ۶.۴: جمع‌بندی نتایج، مهارت‌های اکتسابی و درس‌های آموخته‌شده", "صفحه ۲۷"],
        ["  • ۶.۵ تا ۶.۶: نقشه راه توسعه فازهای آتی و پیشنهادات به کارآموزان", "صفحه ۲۸"],
        ["بخش پایانی: پیوست‌ها و فهرست مراجع علمی", "صفحه ۲۹ تا ۳۰"],
        ["  • پیوست الف و ب: ساختار فایل‌ها، پکیج‌ها و راهنمای راه‌اندازی پروژه", "صفحه ۲۹"],
        ["  • مراجع و استانداردهای علمی، کتاب‌ها، مقالات و منابع رسمی وب", "صفحه ۳۰"]
    ]
    add_table_data(doc, ["سرفصل‌های گزارش کارآموزی", "محدوده صفحات"], toc_data, caption="جدول فهرست تفصیلی فصول و مباحث گزارش")

def build_page_05(doc):
    """صفحه ۵: فصل اول - پیشینه، مأموریت و زیست‌بوم ناحیه نوآوری شریف"""
    add_heading_1(doc, "فصل اول: معرفی محل کارآموزی، زیست‌بوم و اهداف دوره", space_before=6, space_after=8)
    
    add_heading_2(doc, "۱.۱ مقدمه و تاریخچه پیدایش استارتاپ استودیوی شریف کلینیک", space_before=4, space_after=3)
    add_p(doc, "استارتاپ استودیو و شتاب‌دهنده تخصصی «شریف کلینیک» (Sharif Clinic) با درک ضرورت عمیق تحول دیجیتال در نظام سلامت کشور و با هدف ایجاد پلی پایدار میان فناوری‌های نوین مهندسی نرم‌افزار، هوش مصنوعی و تخصص‌های بالینی پزشکی شکل گرفت. در سال‌های اخیر، عدم همگرایی کارآمد میان مهندسان خلاق و کادر پزشکی باعث شده بود که بسیاری از نیازهای مبرم نظام درمانی بدون پاسخ فناورانه باقی بماند. شریف کلینیک به عنوان یک بازیگر پیشرو، فرآیند نوآوری را با رویکرد Venture Studio پایه‌ریزی کرده است تا ایده‌های ناب از مرحله شناسایی بالینی تا بلوغ محصول و نفوذ به بازار هدایت شوند.",
          space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_2(doc, "۱.۲ زیست‌بوم نوآوری دانشگاه صنعتی شریف و استقرار در ناحیه نوآوری", space_before=5, space_after=3)
    add_p(doc, "محل استقرار شریف کلینیک در مجاورت مستقیم دانشگاه صنعتی شریف (واقع در خیابان شهید ولی‌الله صادقی) و در قلب «ناحیه نوآوری شریف» قرار دارد. این موقعیت جغرافیایی و راهبردی فرصتی بی‌بدیل پدید آورده است تا پلتفرم به صورت پیوسته از ظرفیت‌های پژوهشی نخبگان، دانشجویان برتر مهندسی کامپیوتر و آزمایشگاه‌های پیشرفته بهره‌مند گردد. این دوره کارآموزی دانشجوی دانشگاه آزاد اسلامی واحد پرند فرصتی مغتنم جهت پیوند توانمندی‌های دانشگاهی با قطب نوآوری شریف و شرکت‌های دانش‌بنیان مستقر در پایتخت را فراهم ساخت.",
          space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_heading_2(doc, "۱.۳ مأموریت، ارزش‌ها و شعار راهبردی سازمانی", space_before=5, space_after=3)
    add_p(doc, "مأموریت بنیادی شریف کلینیک عبارت است از «کوتاه کردن فاصله میان ایده پزشکی تا نجات جان بیمار به واسطه زیرساخت‌های نرم‌افزاری مقیاس‌پذیر و ایمن». شعار استراتژیک این مجموعه «دقت مهندسی و تخصص پزشکی، آینده سلامت را امروز بسازید» نشان‌دهنده اراده محوری تیم در ادغام معیارهای سخت‌گیرانه مهندسی با اخلاق و موازین سلامت است.",
          space_after=4, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    
    add_callout(doc, "سه رکن بنیادین ارزش‌های سازمانی در شریف کلینیک", [
        "۱. دقت بالینی و امنیت اطلاعات (Clinical Rigor): تعهد خدشه‌ناپذیر به حفظ سلامت بیمار و حریم خصوصی داده‌های بالینی.",
        "۲. چابکی و تعالی مهندسی (Engineering Excellence): به کارگیری مدرن‌ترین الگوهای معماری نرم‌افزار و پایداری کلاینت-سرور.",
        "۳. هم‌افزایی انسان‌محور (Human-Centric Synergy): تبدیل فناوری به ابزاری آرامش‌بخش برای بیمار و تسهیل‌کننده کار پزشک."
    ], bg_hex="F0FDF4", border_color="0D9488")

print("Module pages_01_to_05 loaded.")
