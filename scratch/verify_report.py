import docx
import os
import sys

doc = docx.Document("Sharif_Clinic_Internship_Report.docx")
size = os.path.getsize("Sharif_Clinic_Internship_Report.docx")

page_breaks = 0
for p in doc.paragraphs:
    if any(br.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type") == "page" 
           for run in p.runs for br in run._r.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br")):
        page_breaks += 1

all_text = " ".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for row in t.rows for c in row.cells])
words = len(all_text.split())
chars = len(all_text)

print(f"File Size: {size/1024:.1f} KB")
print(f"Paragraphs: {len(doc.paragraphs)}")
print(f"Tables: {len(doc.tables)}")
print(f"Total Words: {words}")
print(f"Total Characters: {chars}")
print(f"Page breaks count: {page_breaks}")
print(f"Calculated Pages (breaks + 1): {page_breaks + 1}")
