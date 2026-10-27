import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def apply_font_to_run(run, font_name="Vazirmatn", size_pt=11, bold=False, italic=False, color_rgb=(30, 41, 59)):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color_rgb)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="{font_name}" w:hAnsi="{font_name}" w:cs="{font_name}"/>')
    rPr.append(rFonts)
    rtl = parse_xml(f'<w:rtl {nsdecls("w")}/>')
    rPr.append(rtl)

def add_p(doc, text, bold_prefix="", size_pt=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=3.5, space_before=0, color_rgb=(30, 41, 59)):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = align
    p.paragraph_format.bidi = True
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        apply_font_to_run(r_pre, size_pt=size_pt, bold=True, color_rgb=color_rgb)
    if text:
        r_text = p.add_run(text)
        apply_font_to_run(r_text, size_pt=size_pt, bold=False, color_rgb=color_rgb)
    return p

doc = Document()
add_p(doc, "این یک متن آزمایشی است برای بررسی صحت قلم و جهت راست‌چین در مایکروسافت ورد.", bold_prefix="نکته مهم: ")
doc.save("scratch/test_text_flow.docx")
print("Text flow test passed!")
