# -*- coding: utf-8 -*-
"""
Master Builder Script for 30-Page Sharif Clinic Internship Report.
Generates Sharif_Clinic_Internship_Report.docx with high visual excellence,
academic rigor, Persian RTL typography, and clean pagination.
"""

import os
import sys

# Ensure both current directory and scratch directory are in sys.path
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.dirname(BASE_DIR)
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from report_helpers import apply_run_font
from pages_01_to_05 import build_page_01, build_page_02, build_page_03, build_page_04, build_page_05
from pages_06_to_10 import build_page_06, build_page_07, build_page_08, build_page_09, build_page_10
from pages_11_to_15 import build_page_11, build_page_12, build_page_13, build_page_14, build_page_15
from pages_16_to_20 import build_page_16, build_page_17, build_page_18, build_page_19, build_page_20
from pages_21_to_25 import build_page_21, build_page_22, build_page_23, build_page_24, build_page_25
from pages_26_to_30 import build_page_26, build_page_27, build_page_28, build_page_29, build_page_30

def create_report(output_path):
    print(f"Initializing Document for {output_path}...")
    doc = Document()
    
    # Configure Page Setup (A4)
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    
    # Configure Header & Footer
    section.different_first_page_header_footer = True
    
    # First page header & footer empty
    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""
    first_footer = section.first_page_footer
    first_footer.paragraphs[0].text = ""
    
    # Standard Header (pages 2+)
    header = section.header
    p_h = header.paragraphs[0]
    p_h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_h.paragraph_format.bidi = True
    p_h.paragraph_format.space_after = Pt(2)
    r_h = p_h.add_run("گزارش تخصصی کارآموزی مهندسی کامپیوتر | پلتفرم سلامت هوشمند شریف کلینیک (sharifclinic.co)")
    apply_run_font(r_h, size_pt=8.5, color_rgb=(100, 116, 139))
    
    # Subtle bottom border on header
    pPr = p_h._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="2" w:color="CBD5E1"/></w:pBdr>')
    pPr.append(pBdr)
    
    # Standard Footer (pages 2+)
    footer = section.footer
    p_f = footer.paragraphs[0]
    p_f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_f.paragraph_format.bidi = True
    p_f.paragraph_format.space_before = Pt(4)
    
    r_f_txt = p_f.add_run("دانشگاه صنعتی شریف - دانشکده مهندسی کامپیوتر  |  صفحه ")
    apply_run_font(r_f_txt, size_pt=9, color_rgb=(100, 116, 139))
    
    r_page = p_f.add_run()
    apply_run_font(r_page, size_pt=9, bold=True, color_rgb=(15, 53, 98))
    fld = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    r_page._r.append(fld)
    
    r_f_total = p_f.add_run(" از ۳۰")
    apply_run_font(r_f_total, size_pt=9, color_rgb=(100, 116, 139))
    
    # Sequentially build all 30 pages
    pages = [
        build_page_01, build_page_02, build_page_03, build_page_04, build_page_05,
        build_page_06, build_page_07, build_page_08, build_page_09, build_page_10,
        build_page_11, build_page_12, build_page_13, build_page_14, build_page_15,
        build_page_16, build_page_17, build_page_18, build_page_19, build_page_20,
        build_page_21, build_page_22, build_page_23, build_page_24, build_page_25,
        build_page_26, build_page_27, build_page_28, build_page_29, build_page_30
    ]
    
    for i, page_fn in enumerate(pages):
        page_num = i + 1
        print(f"Building Page {page_num:02d} / 30...")
        page_fn(doc)
        if page_num < 30:
            doc.add_page_break()
            
    print(f"Saving final report to {output_path}...")
    doc.save(output_path)
    print("Report generated and saved successfully!")

if __name__ == "__main__":
    out_file = "Sharif_Clinic_Internship_Report.docx"
    create_report(out_file)
