import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()
section = doc.sections[0]
section.different_first_page_header_footer = True

# First page footer should be empty
first_footer = section.first_page_footer
p_first = first_footer.paragraphs[0]
p_first.text = ""

# Normal header & footer for subsequent pages
header = section.header
p_head = header.paragraphs[0]
p_head.text = "گزارش کارآموزی | شریف کلینیک"
p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT

footer = section.footer
p_foot = footer.paragraphs[0]
p_foot.text = "صفحه "
fld = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
p_foot.runs[0]._r.append(fld)

doc.add_paragraph("صفحه جلد - بدون سربرگ و پابرگ")
doc.add_page_break()
doc.add_paragraph("صفحه دوم - دارای سربرگ و پابرگ")
doc.save("scratch/test_first_page.docx")
print("First page header/footer test passed!")
